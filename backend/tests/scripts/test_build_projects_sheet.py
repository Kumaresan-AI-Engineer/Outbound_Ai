import importlib.util
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from app.domain.projects.excel_import import parse_projects_workbook

SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "build_projects_sheet.py"
spec = importlib.util.spec_from_file_location("build_projects_sheet", SCRIPT)
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


def _profile(path: Path, paragraphs: list[str]) -> Path:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    document.save(path)
    return path


FITAI = [
    "Jane Doe",
    "jane@example.com | +1 555 0100",  # personal details that must never reach the sheet
    "Project Name: FitAI",
    "Technologies: Python, FastAPI, AWS (S3, SQS), React.",
    "Responsibilities:",
    "API Design: built the APIs.",
    "Description:",
    "FitAI is a meal platform. It uses vector search. A third sentence that should not be in the summary.",
]


def test_parse_projects_reads_name_technologies_and_description(tmp_path):
    lines = builder.read_docx_lines(_profile(tmp_path / "a.docx", FITAI))

    (project,) = builder.parse_projects(lines, "a.docx")

    assert project.name == "FitAI"
    assert project.technologies == ["Python", "FastAPI", "AWS (S3, SQS)", "React"]
    assert project.description.startswith("FitAI is a meal platform.")


def test_technologies_that_continue_on_the_next_line_are_joined(tmp_path):
    lines = [
        "Project Name: X",
        "Technologies: Python, FastAPI,",
        "Chroma DB, Web Scraping.",
        "Description: Does x.",
    ]

    (project,) = builder.parse_projects(lines, "a.docx")

    assert project.technologies == ["Python", "FastAPI", "Chroma DB", "Web Scraping"]


def test_description_on_the_same_line_as_its_label_is_read():
    (project,) = builder.parse_projects(
        ["Project Name: X", "Technologies: Go", "Description: One. Two."], "a.docx"
    )

    assert project.description == "One. Two."


def test_multiple_projects_in_one_document_are_split_correctly():
    lines = [
        "Project Name: A",
        "Technologies: Go",
        "Description: About A.",
        "Project Name: B",
        "Technologies: Rust",
        "Description: About B.",
    ]

    projects = builder.parse_projects(lines, "a.docx")

    assert [(p.name, p.technologies, p.description) for p in projects] == [
        ("A", ["Go"], "About A."),
        ("B", ["Rust"], "About B."),
    ]


def test_collect_keeps_the_richest_variant_across_files_and_recurses(tmp_path):
    _profile(tmp_path / "short.docx", ["Project Name: FitAi", "Technologies: Python", "Description: Short."])
    nested = tmp_path / "Role Profiles"
    nested.mkdir()
    _profile(nested / "long.docx", FITAI)

    (project,) = builder.collect_projects(tmp_path)

    assert project.name == "FitAI" and project.source == "long.docx"


def test_collect_ignores_word_lock_files_and_unreadable_documents(tmp_path, capsys):
    _profile(tmp_path / "good.docx", FITAI)
    (tmp_path / "~$good.docx").write_bytes(b"lock")
    (tmp_path / "broken.docx").write_bytes(b"this is not a zip")

    projects = builder.collect_projects(tmp_path)

    assert [p.name for p in projects] == ["FitAI"]
    assert "broken.docx" in capsys.readouterr().err


def test_summary_is_the_first_two_sentences():
    text = "One. Two! Three? Four."

    assert builder.summarize(text) == "One. Two!"


def test_long_summary_is_trimmed_at_a_word_boundary():
    summary = builder.summarize("word " * 200)

    assert len(summary) <= 400 and summary.endswith("...")


def test_clean_normalises_typography():
    assert builder.clean("It\u2019s \u201cgood\u201d \u2013 really\u00a0so") == 'It\'s "good" - really so'


def test_technology_grouping():
    groups = builder.group_technologies(
        [
            "React.js",
            "FastAPI",
            "Milvus (Vector DB)",
            "Azure (Blob Storage, App Services)",
            "Twilio",
            "AWS S3",
        ]
    )

    assert groups == {
        "frontend": ["React.js"],
        "backend": ["FastAPI"],
        "database": ["Milvus (Vector DB)"],
        "cloud_devops": ["Azure (Blob Storage, App Services)", "AWS S3"],
    }  # Twilio is in no group; the importer keeps it under "other"


def test_personal_details_never_reach_the_sheet(tmp_path):
    _profile(tmp_path / "a.docx", FITAI)
    target = tmp_path / "out.xlsx"

    builder.write_workbook(builder.collect_projects(tmp_path), target)

    cells = " ".join(str(c.value) for row in load_workbook(target).active.iter_rows() for c in row if c.value)
    assert "jane@example.com" not in cells and "Jane Doe" not in cells and "555 0100" not in cells


def test_curated_and_uncurated_projects(tmp_path):
    _profile(
        tmp_path / "a.docx",
        FITAI + ["Project Name: Brand New", "Technologies: Go", "Description: New thing."],
    )

    rows = {r[0]: r for r in map(builder.sheet_row, builder.collect_projects(tmp_path))}

    assert (
        rows["FitAI"][1] == "Travel & Hospitality" and rows["FitAI"][4]
    )  # curated domain + business problem
    assert rows["Brand New"][1] == "Other" and rows["Brand New"][4] == ""  # falls back, still importable


def test_generated_sheet_round_trips_through_the_importer(tmp_path):
    _profile(tmp_path / "a.docx", FITAI)
    target = tmp_path / "out.xlsx"
    builder.write_workbook(builder.collect_projects(tmp_path), target)

    result = parse_projects_workbook(target.read_bytes())

    assert result.errors == [] and result.warnings == []
    meta = result.projects[0].metadata
    assert meta.title == "FitAI" and meta.domain == "Travel & Hospitality"
    assert meta.related_domains == ["Healthcare"]
    assert meta.tech_stack.frontend == ["React"] and meta.tech_stack.backend == ["Python", "FastAPI"]
    assert meta.tech_stack.cloud_devops == ["AWS (S3, SQS)"]
    assert len(meta.key_features) >= 3 and meta.ai_ml_components


def test_committed_sample_has_every_curated_project_and_no_uncurated_gaps():
    result = parse_projects_workbook((SCRIPT.parent / "sample_projects.xlsx").read_bytes())

    assert result.errors == [] and result.warnings == []
    assert {p.name for p in result.projects} == {
        "DK Tire B2B Platform",
        "Eagle-Bot (SAAS)",
        "FitAI",
        "ImpactX",
        "NeonAI Co-pilot Agents",
        "Orion-Chat",
        "Perky Pet",
        "Quadrix AI",
        "SoftSuave - AI Assistant",
        "Tireplanet",
        "Zysec AI",
    }
    for project in result.projects:
        assert project.metadata.business_problem and project.metadata.solution_provided, project.name


def test_cli_rejects_a_folder_that_does_not_exist(monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", ["build_projects_sheet.py", "Z:/definitely/not/here"])

    assert builder.main() == 1
    assert "Not a folder" in capsys.readouterr().err
