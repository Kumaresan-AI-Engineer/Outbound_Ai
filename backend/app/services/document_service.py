"""Text extraction for uploaded project documents (PDF / Word).

All extractors are synchronous (pypdf and python-docx are blocking) — callers
run them via asyncio.to_thread so the event loop stays free.
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".doc", ".docx"}
MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB


class ExtractionError(Exception):
    """Raised when no usable text could be extracted from a document."""


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(file_path: Path) -> str:
    from docx import Document

    doc = Document(str(file_path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables often hold the interesting facts (tech stack, team, duration)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(file_path: str, file_type: str) -> str:
    """Extract plain text from a document. Raises ExtractionError on failure."""
    path = Path(file_path)
    if not path.exists():
        raise ExtractionError("Uploaded file is missing on disk.")

    try:
        if file_type == "pdf":
            text = _extract_pdf(path)
        elif file_type == "docx":
            text = _extract_docx(path)
        elif file_type == "doc":
            # Some .doc files are actually docx in disguise; try that first.
            try:
                text = _extract_docx(path)
            except Exception:
                raise ExtractionError(
                    "Legacy .doc format could not be parsed. "
                    "Please re-save the document as .docx or PDF and upload again."
                )
        else:
            raise ExtractionError(f"Unsupported file type: {file_type}")
    except ExtractionError:
        raise
    except Exception as e:
        logger.error(f"[EXTRACT] Failed for {path.name}: {e}")
        raise ExtractionError(f"Could not read the document: {e}")

    if len(text.strip()) < 50:
        raise ExtractionError(
            "The document contains no readable text (it may be scanned images). "
            "Please upload a text-based PDF or Word document."
        )
    return text
