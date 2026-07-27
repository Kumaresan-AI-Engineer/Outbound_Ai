import io
import re
import asyncio
import logging
from datetime import datetime
from fastapi import APIRouter, HTTPException, UploadFile, File
from bson import ObjectId
from app.database import clients_collection, contacts_collection
from app.models.schemas import ClientResponse, ClientUploadResult
from app.services.client_agent import enrich_client

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/clients", tags=["clients"])

MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
MAX_ROWS = 1000

# Accepted header spellings, normalized to lowercase alphanumerics.
# The sheet must contain one match for each required field.
HEADER_ALIASES = {
    "name": {"clientname", "name", "client", "fullname"},
    "company": {"companydetails", "company", "companyname", "organisation", "organization"},
    "project": {"project", "clientproject", "projectdetails", "projectname"},
    "phone": {"contactnumber", "phone", "phonenumber", "contact", "mobile", "mobilenumber"},
    "secondary_phone": {"secondaryphone", "secondarynumber", "alternatephone", "alternatenumber", "phone2", "contactnumber2"},
}
REQUIRED_FIELDS = {"name", "phone"}

# Two-sheet "Lead format" workbook (Apollo-style lead-gen export): a
# "Contacts format" sheet (one row per person) and a "Lead list format"
# sheet (one row per company/campaign) joined on company name.
CONTACT_FIELD_ALIASES = {
    "first_name": {"firstname"},
    "last_name": {"lastname"},
    "phone": {"numbercontactnumbers", "number", "contactnumber", "contactnumbers", "phone"},
    "secondary_phone": {"secondaryphone", "secondarynumber", "alternatephone", "alternatenumber", "phone2"},
    "link_name": {"linknamelinks", "linkname", "company", "companyname"},
}
LEAD_FIELD_ALIASES = {
    "company": {"companyname", "company"},
    "industries": {"industries", "industry"},
}

# Cells that pack more than one number together (the Lead format's "Number
# (Contact Numbers)" column is plural for a reason) get split on these.
_PHONE_SPLIT_RE = re.compile(r"[;,\n]+")


def _normalize_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _normalize_phone(value) -> str:
    s = str(value or "").strip()
    # Excel often stores numbers as floats ("9876543210.0")
    if s.endswith(".0"):
        s = s[:-2]
    digits = re.sub(r"[^\d+]", "", s)
    # Numbers without a country-code prefix (e.g. "49 30 6322265708") aren't
    # dialable via Twilio as-is; assume the leading digits already are the
    # country code and just add the "+" E.164 needs.
    if digits and not digits.startswith("+") and len(digits) >= 10:
        digits = "+" + digits
    return digits


def _split_phones(value) -> list:
    s = str(value or "").strip()
    if not s:
        return []
    parts = [_normalize_phone(p) for p in _PHONE_SPLIT_RE.split(s) if p.strip()]
    return [p for p in parts if p]


def _extract_phones(values_by_field: dict) -> tuple:
    """Returns (primary, secondary) phone numbers for one row. An explicit
    secondary-phone column always wins; otherwise a second number packed
    into the primary phone cell (e.g. "123, 456") is used as the fallback."""
    phones = _split_phones(values_by_field.get("phone"))
    primary = phones[0] if phones else ""
    secondary = _normalize_phone(values_by_field.get("secondary_phone")) or (phones[1] if len(phones) > 1 else "")
    return primary, secondary


def _map_headers(header_row, aliases: dict) -> dict:
    """Map column index -> field name using the given alias table."""
    mapping = {}
    for idx, cell in enumerate(header_row):
        normalized = _normalize_header(cell)
        for field, field_aliases in aliases.items():
            if normalized in field_aliases and field not in mapping.values():
                mapping[idx] = field
    return mapping


def _find_sheet_by_keyword(workbook, keyword: str):
    for name in workbook.sheetnames:
        if keyword in _normalize_header(name):
            return workbook[name]
    return None


def _extract_lead_format_rows(workbook, contact_rows: list) -> list:
    """Parse the two-sheet Lead format into the common
    {name, company, project, phone, secondary_phone} row shape the rest of
    the importer already understands. "project" carries the Lead sheet's
    Industries value through unchanged so it flows into the existing AI
    domain classifier exactly like a manual Project entry would."""
    lead_sheet = _find_sheet_by_keyword(workbook, "lead")
    lead_by_company = {}
    fallback_industries = ""
    if lead_sheet is not None:
        lead_rows = list(lead_sheet.iter_rows(values_only=True))
        if lead_rows:
            lead_header_map = _map_headers(lead_rows[0], LEAD_FIELD_ALIASES)
            for raw in lead_rows[1:]:
                lead = {}
                for idx, field in lead_header_map.items():
                    value = raw[idx] if idx < len(raw) else None
                    lead[field] = str(value or "").strip()
                company = lead.get("company", "")
                industries = lead.get("industries", "")
                if company:
                    lead_by_company[_normalize_header(company)] = {
                        "company": company,
                        "industries": industries,
                    }
                if industries and not fallback_industries:
                    # Real-world exports often leave "Company Name" blank on
                    # every lead row, so an exact join can match nothing -
                    # fall back to the first known industry as a shared
                    # campaign hint rather than dropping the signal entirely.
                    fallback_industries = industries

    contact_header_map = _map_headers(contact_rows[0], CONTACT_FIELD_ALIASES)
    data_rows = []
    for raw in contact_rows[1:MAX_ROWS + 1]:
        values_by_field = {}
        for idx, field in contact_header_map.items():
            values_by_field[field] = raw[idx] if idx < len(raw) else None
        primary, secondary = _extract_phones(values_by_field)

        link_name = str(values_by_field.get("link_name") or "").strip()
        matched = lead_by_company.get(_normalize_header(link_name)) if link_name else None
        if matched:
            company = matched["company"] or link_name
            industries = matched["industries"]
        else:
            company = link_name
            industries = fallback_industries

        first_name = str(values_by_field.get("first_name") or "").strip()
        last_name = str(values_by_field.get("last_name") or "").strip()
        data_rows.append({
            "name": f"{first_name} {last_name}".strip(),
            "company": company,
            "project": industries,
            "phone": primary,
            "secondary_phone": secondary,
        })
    return data_rows


def serialize_client(doc) -> dict:
    return {
        "id": str(doc["_id"]),
        "name": doc.get("name", ""),
        "company": doc.get("company", ""),
        "project": doc.get("project", ""),
        "phone": doc.get("phone", ""),
        "contact_id": doc.get("contact_id"),
        "domain": doc.get("domain"),
        "related_domains": doc.get("related_domains", []),
        "matched_project_ids": doc.get("matched_project_ids", []),
        "enrichment_status": doc.get("enrichment_status", "pending"),
        "source_file": doc.get("source_file", ""),
        "created_at": doc.get("created_at", datetime.utcnow()),
    }


async def _upsert_contact(row: dict) -> str:
    """Create or update the callable contact for an imported client row.
    Existing contacts keep their data — we only fill blanks."""
    existing = await contacts_collection.find_one({"phone": row["phone"]})
    note = f"Client project: {row['project']}" if row["project"] else ""
    secondary_phone = row.get("secondary_phone", "")
    if existing:
        update = {}
        if not existing.get("company") and row["company"]:
            update["company"] = row["company"]
        if not existing.get("notes") and note:
            update["notes"] = note
        if not existing.get("secondary_phone") and secondary_phone:
            update["secondary_phone"] = secondary_phone
        if update:
            await contacts_collection.update_one({"_id": existing["_id"]}, {"$set": update})
        return str(existing["_id"])

    result = await contacts_collection.insert_one({
        "name": row["name"],
        "phone": row["phone"],
        "secondary_phone": secondary_phone,
        "company": row["company"],
        "status": "new",
        "notes": note,
        "last_called": None,
        "created_at": datetime.utcnow(),
    })
    return str(result.inserted_id)


@router.post("/upload", response_model=ClientUploadResult)
async def upload_clients(file: UploadFile = File(...)):
    """Import clients from an Excel (.xlsx) file. Each valid row is stored as
    a client, upserted into contacts (so it's immediately callable), and
    queued for AI domain classification + project matching."""
    filename = file.filename or ""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "xls":
        raise HTTPException(
            status_code=400,
            detail="Legacy .xls format is not supported. Please save the file as .xlsx and try again.",
        )
    if ext not in ("xlsx", "xlsm"):
        raise HTTPException(status_code=400, detail="Please upload an Excel file (.xlsx).")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="File too large. Maximum size is 5 MB.")

    try:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        logger.error(f"[CLIENTS] Failed to parse Excel '{filename}': {e}")
        raise HTTPException(status_code=400, detail="Could not read the Excel file. Is it a valid .xlsx?")

    # Detect the two-sheet Lead format (a "Contacts format" sheet whose
    # header maps a phone column) before falling back to the original
    # single-sheet Client/Company/Project/Contact layout.
    data_rows = None
    contacts_sheet = _find_sheet_by_keyword(workbook, "contact")
    if contacts_sheet is not None:
        contact_rows = list(contacts_sheet.iter_rows(values_only=True))
        if contact_rows and "phone" in _map_headers(contact_rows[0], CONTACT_FIELD_ALIASES).values():
            data_rows = _extract_lead_format_rows(workbook, contact_rows)

    if data_rows is None:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise HTTPException(status_code=400, detail="The Excel sheet is empty.")

        header_map = _map_headers(rows[0], HEADER_ALIASES)
        mapped_fields = set(header_map.values())
        missing = REQUIRED_FIELDS - mapped_fields
        if missing:
            pretty = {"name": "Client Name", "phone": "Contact Number"}
            raise HTTPException(
                status_code=400,
                detail=f"Missing required column(s): {', '.join(pretty[f] for f in sorted(missing))}. "
                       "Expected columns: Client Name, Company Details, Project, Contact Number.",
            )

        data_rows = []
        for raw in rows[1:MAX_ROWS + 1]:
            values_by_field = {}
            for idx, field in header_map.items():
                values_by_field[field] = raw[idx] if idx < len(raw) else None
            primary, secondary = _extract_phones(values_by_field)
            data_rows.append({
                "name": str(values_by_field.get("name") or "").strip(),
                "company": str(values_by_field.get("company") or "").strip(),
                "project": str(values_by_field.get("project") or "").strip(),
                "phone": primary,
                "secondary_phone": secondary,
            })

    imported = updated = skipped = 0
    errors = []
    new_client_ids = []

    for i, row in enumerate(data_rows, start=2):  # start=2 → Excel row numbers
        if not any(row.values()):
            continue  # blank row
        if not row["name"]:
            skipped += 1
            errors.append({"row": i, "message": "Missing client name"})
            continue
        if not row["phone"] or len(re.sub(r"\D", "", row["phone"])) < 7:
            skipped += 1
            errors.append({"row": i, "message": f"Invalid or missing contact number for '{row['name']}'"})
            continue

        try:
            contact_id = await _upsert_contact(row)
            existing = await clients_collection.find_one({"phone": row["phone"], "project": row["project"]})
            if existing:
                await clients_collection.update_one(
                    {"_id": existing["_id"]},
                    {"$set": {
                        "name": row["name"],
                        "company": row["company"],
                        "contact_id": contact_id,
                        "source_file": filename,
                        "updated_at": datetime.utcnow(),
                    }},
                )
                new_client_ids.append(str(existing["_id"]))
                updated += 1
            else:
                result = await clients_collection.insert_one({
                    **row,
                    "contact_id": contact_id,
                    "domain": None,
                    "related_domains": [],
                    "matched_project_ids": [],
                    "enrichment_status": "pending",
                    "source_file": filename,
                    "created_at": datetime.utcnow(),
                    "updated_at": datetime.utcnow(),
                })
                new_client_ids.append(str(result.inserted_id))
                imported += 1
        except Exception as e:
            logger.error(f"[CLIENTS] Row {i} failed: {e}")
            skipped += 1
            errors.append({"row": i, "message": "Could not save this row"})

    # Queue AI domain classification + project matching in the background
    for client_id in new_client_ids:
        asyncio.create_task(enrich_client(client_id))

    logger.info(
        f"[CLIENTS] Imported '{filename}': {imported} new, {updated} updated, {skipped} skipped"
    )
    return {
        "total_rows": imported + updated + skipped,
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors[:20],
    }


@router.get("/", response_model=list[ClientResponse])
async def get_clients():
    clients = []
    async for doc in clients_collection.find().sort("created_at", -1):
        clients.append(serialize_client(doc))
    return clients


@router.delete("/{client_id}")
async def delete_client(client_id: str):
    if not ObjectId.is_valid(client_id):
        raise HTTPException(status_code=404, detail="Client not found")
    result = await clients_collection.delete_one({"_id": ObjectId(client_id)})
    if result.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Client not found")
    return {"message": "Client deleted"}
