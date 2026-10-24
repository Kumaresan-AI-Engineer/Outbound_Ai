"""Excel import of the project knowledge base.

A sheet row becomes a finished project record - the same `metadata` shape the
AI document pipeline produces - so no model call is involved and the result is
exactly what the sheet says. Imported projects feed client matching and the
live-call coaching the same way uploaded documents do.
"""

import asyncio
import io
import logging
import re
from dataclasses import dataclass, field
from datetime import datetime

from app.ai.agents.client_domain_agent import refresh_matches_for_project
from app.ai.schemas.project_extraction import CANONICAL_DOMAINS, ProjectMetadata, TechStack
from app.domain.spreadsheet import map_headers
from app.repositories.projects_repo import projects_repo

logger = logging.getLogger(__name__)

MAX_ROWS = 500
SOURCE = "excel"

HEADER_ALIASES = {
    "name": {"projectname", "name", "project", "title"},
    "domain": {"domain", "industry", "vertical"},
    "related_domains": {"relateddomains", "otherdomains", "secondarydomains"},
    "summary": {"summary", "description", "overview"},
    "business_problem": {"businessproblem", "problem", "challenge"},
    "solution_provided": {"solutionprovided", "solution"},
    "key_features": {"keyfeatures", "features"},
    "technologies": {"technologies", "technology", "techstack", "tools"},
    "ai_ml_components": {"aimlcomponents", "aicomponents", "aiml", "mlcomponents"},
    "frontend": {"frontend"},
    "backend": {"backend"},
    "database": {"database", "databases"},
    "cloud_devops": {"clouddevops", "cloud", "devops"},
}
REQUIRED_FIELDS = {"name"}
EXPECTED_COLUMNS = (
    "Required: Project Name. Recommended: Domain, Summary, Technologies. Optional: Related Domains, "
    "Business Problem, Solution Provided, Key Features, AI/ML Components, Frontend, Backend, Database, "
    "Cloud / DevOps."
)

_TECH_GROUPS = ("frontend", "backend", "database", "cloud_devops")


class SheetFormatError(ValueError):
    """The file as a whole can't be imported (unreadable, empty, wrong columns)."""


@dataclass
class RowIssue:
    row: int
    message: str


@dataclass
class ParsedProject:
    row: int
    name: str
    metadata: ProjectMetadata


@dataclass
class ParseResult:
    projects: list[ParsedProject] = field(default_factory=list)
    errors: list[RowIssue] = field(default_factory=list)
    warnings: list[RowIssue] = field(default_factory=list)
    total_rows: int = 0


def _domain_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower().replace("&", "and"))


_CANONICAL_BY_KEY = {_domain_key(d): d for d in CANONICAL_DOMAINS}


def normalize_domain(value: str) -> str | None:
    """Canonical domain for a free-text value ('retail and e-commerce' ->
    'Retail & E-Commerce'), or None when it isn't one we know."""
    return _CANONICAL_BY_KEY.get(_domain_key(value or ""))


def split_list(value, separators: str = ",;\n") -> list[str]:
    """Split a cell into items. Separators inside parentheses are kept, so
    'AWS (S3, SQS, Lambda), Redis' is two items, not four."""
    items: list[str] = []
    current: list[str] = []
    depth = 0
    for char in str(value or ""):
        if char in "([":
            depth += 1
        elif char in ")]":
            depth = max(0, depth - 1)
        if char in separators and depth == 0:
            items.append("".join(current))
            current = []
        else:
            current.append(char)
    items.append("".join(current))
    cleaned = (" ".join(i.split()).strip(" .") for i in items)
    return list(dict.fromkeys(i for i in cleaned if i))  # drop blanks, keep first of any duplicate


def _cell(values: dict, key: str) -> str:
    value = values.get(key)
    return "" if value is None else str(value).strip()


def _build_tech(values: dict) -> tuple[list[str], TechStack]:
    technologies = split_list(_cell(values, "technologies"))
    groups = {name: split_list(_cell(values, name)) for name in _TECH_GROUPS}
    if not any(groups.values()):
        return technologies, TechStack()
    grouped = {t.lower() for items in groups.values() for t in items}
    # The UI shows the grouped view when it exists, so anything left out of every
    # group goes under "other" rather than silently disappearing.
    leftover = [t for t in technologies if t.lower() not in grouped]
    all_tech = list(dict.fromkeys([*technologies, *(t for items in groups.values() for t in items)]))
    return all_tech, TechStack(**groups, other=leftover)


def _parse_row(row_number: int, values: dict, result: ParseResult) -> None:
    name = _cell(values, "name")
    if not name:
        result.errors.append(RowIssue(row_number, "Missing project name"))
        return
    if not (_cell(values, "summary") or _cell(values, "technologies")):
        result.errors.append(RowIssue(row_number, f"'{name}' needs at least a Summary or Technologies"))
        return

    raw_domain = _cell(values, "domain")
    domain = normalize_domain(raw_domain) if raw_domain else "Other"
    if domain is None:
        result.warnings.append(
            RowIssue(row_number, f"Unknown domain '{raw_domain}' for '{name}' - saved as Other")
        )
        domain = "Other"

    related: list[str] = []
    for item in split_list(_cell(values, "related_domains")):
        canonical = normalize_domain(item)
        if canonical:
            related.append(canonical)
        else:
            result.warnings.append(
                RowIssue(row_number, f"Unknown related domain '{item}' for '{name}' - ignored")
            )

    technologies, tech_stack = _build_tech(values)
    metadata = ProjectMetadata(
        title=name,
        domain=domain,
        related_domains=related,
        technologies=technologies,
        summary=_cell(values, "summary"),
        key_features=split_list(_cell(values, "key_features"), separators=";\n"),
        business_problem=_cell(values, "business_problem"),
        solution_provided=_cell(values, "solution_provided"),
        ai_ml_components=split_list(_cell(values, "ai_ml_components"), separators=";\n"),
        tech_stack=tech_stack.model_dump(),
    )
    result.projects.append(ParsedProject(row_number, name, metadata))


def parse_projects_workbook(content: bytes) -> ParseResult:
    """Read the first sheet into validated project metadata. Row problems are
    collected, not raised; only a file-level problem raises SheetFormatError."""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
    except Exception as e:
        raise SheetFormatError("Could not read the Excel file. Is it a valid .xlsx?") from e
    if not rows:
        raise SheetFormatError("The Excel sheet is empty.")

    header_map = map_headers(rows[0], HEADER_ALIASES)
    if REQUIRED_FIELDS - set(header_map.values()):
        raise SheetFormatError(f"Missing required column: Project Name. {EXPECTED_COLUMNS}")

    result = ParseResult()
    for row_number, raw in enumerate(rows[1 : MAX_ROWS + 1], start=2):  # start=2 -> Excel row numbers
        values = {field_name: raw[idx] if idx < len(raw) else None for idx, field_name in header_map.items()}
        if not any(str(v).strip() for v in values.values() if v is not None):
            continue  # blank row
        result.total_rows += 1
        _parse_row(row_number, values, result)
    return result


async def _find_by_name(name: str) -> dict | None:
    return await projects_repo.find_one({"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}})


async def _refresh_matches(project_ids: list[str]) -> None:
    for project_id in project_ids:
        try:
            await refresh_matches_for_project(project_id)
        except Exception as e:
            logger.error(f"[PROJECTS] Match refresh failed for {project_id}: {type(e).__name__}")


async def import_projects(content: bytes, filename: str) -> dict:
    """Parse the workbook and store its projects. Re-importing the same sheet
    updates the Excel-sourced projects in place; a name that already belongs to
    an uploaded document is skipped rather than overwritten."""
    parsed = parse_projects_workbook(content)
    errors = [RowIssue(i.row, i.message) for i in parsed.errors]
    imported = updated = 0
    changed_ids: list[str] = []
    now = datetime.utcnow()

    for project in parsed.projects:
        existing = await _find_by_name(project.name)
        if existing and existing.get("source") != SOURCE:
            errors.append(
                RowIssue(
                    project.row,
                    f"'{project.name}' already exists from an uploaded document - not overwritten",
                )
            )
            continue
        fields = {
            "name": project.name,
            "file_name": filename,
            "file_type": "xlsx",
            "processing_status": "completed",
            "processing_error": None,
            "metadata": project.metadata.model_dump(),
            "updated_at": now,
        }
        if existing:
            await projects_repo.update_one({"_id": existing["_id"]}, {"$set": fields})
            updated += 1
            changed_ids.append(str(existing["_id"]))
        else:
            result = await projects_repo.insert_one(
                {
                    **fields,
                    "file_size": 0,
                    "file_path": "",
                    "extracted_text": "",
                    "source": SOURCE,
                    "created_at": now,
                }
            )
            imported += 1
            changed_ids.append(str(result.inserted_id))

    if changed_ids:
        asyncio.create_task(_refresh_matches(changed_ids))
    logger.info(
        f"[PROJECTS] Excel import '{filename}': {imported} new, {updated} updated, {len(errors)} skipped"
    )
    errors.sort(key=lambda i: i.row)
    return {
        "total_rows": parsed.total_rows,
        "imported": imported,
        "updated": updated,
        "skipped": len(errors),
        "errors": [{"row": i.row, "message": i.message} for i in errors[:20]],
        "warnings": [{"row": i.row, "message": i.message} for i in parsed.warnings[:20]],
    }
